
from pathlib import Path

import pytest
from docx import Document

from edocs_ai.pipeline.steps.extract import extract_document


def test_extract_pdf_text():
    pdf_path = Path("tests/fixtures/sample.pdf")

    result = extract_document(pdf_path)

    assert result.page_count > 0
    assert len(result.pages) == result.page_count
    assert result.source == "sample.pdf"
    assert isinstance(result.full_text, str)
    assert isinstance(result.needs_ocr, bool)


def test_extract_txt(tmp_path):
    txt_path = tmp_path / "sample.txt"
    txt_path.write_text("This is a test document.", encoding="utf-8")

    result = extract_document(txt_path)

    assert result.source == "sample.txt"
    assert "This is a test document." in result.full_text
    assert result.file_type == ".txt"


def test_extract_docx(tmp_path):
    docx_path = tmp_path / "sample.docx"

    document = Document()
    document.add_paragraph("This is a Word document.")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Trevor"

    document.save(docx_path)

    result = extract_document(docx_path)

    assert result.source == "sample.docx"
    assert "This is a Word document." in result.full_text
    assert "Trevor" in result.full_text
    assert result.file_type == ".docx"


def test_unsupported_file_format(tmp_path):
    file_path = tmp_path / "sample.xlsx"
    file_path.write_text("test", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported document format"):
        extract_document(file_path)


def test_missing_file():
    with pytest.raises(FileNotFoundError, match="Document not found"):
        extract_document("does-not-exist.pdf")


def test_extract_pdf_uses_ocr_when_page_has_no_text(tmp_path, monkeypatch):
    import pymupdf

    from edocs_ai.pipeline.steps.extractors import pdf as pdf_extractor

    pdf_path = tmp_path / "scanned.pdf"

    # Create a PDF page with no selectable text.
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.save(pdf_path)

    # Simulate the text returned by Tesseract.
    monkeypatch.setattr(
        pdf_extractor,
        "_run_ocr",
        lambda page: "Recognized text from scanned page.",
    )

    result = extract_document(pdf_path)

    assert result.page_count == 1
    assert result.pages[0].page_number == 1
    assert "Recognized text from scanned page." in result.full_text
    assert result.needs_ocr is False
