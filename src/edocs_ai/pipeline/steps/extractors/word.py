
from pathlib import Path

from docx import Document

from edocs_ai.pipeline.steps.extractors.pdf import (
    ExtractedPage,
    ExtractionResult,
)


def extract_docx(file_path: str | Path) -> ExtractionResult:
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Word document not found: {path}")

    document = Document(path)
    sections: list[str] = []

    # Extract paragraphs.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    # Extract tables without losing row/cell relationships.
    for table_index, table in enumerate(document.tables, start=1):
        sections.append(f"Table {table_index}")

        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            sections.append(" | ".join(cells))

    full_text = "\n\n".join(sections)

    warnings = []
    if not full_text.strip():
        warnings.append("No extractable text found in the Word document.")

    # Word documents don't have reliable fixed page boundaries in this parser.
    return ExtractionResult(
        source=path.name,
        file_type=".docx",
        page_count=0,
        pages=[ExtractedPage(page_number=0, text=full_text)],
        full_text=full_text,
        needs_ocr=False,
        warnings=warnings,
    )
